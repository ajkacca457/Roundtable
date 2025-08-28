import os
from docx import Document  # pip install python-docx
from vector_store import vector_store

DOCS_DIR = os.path.join(os.path.dirname(__file__), "docs")  # points to server/apps/docs

def extract_text_from_docx(path: str) -> str:
    """Extract paragraphs and table text from a .docx file."""
    doc = Document(path)
    parts = []

    # paragraphs
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            parts.append(t)

    # tables (some convos may be in tables)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                t = cell.text.strip()
                if t:
                    parts.append(t)

    return "\n".join(parts)

def main():
    if not os.path.exists(DOCS_DIR):
        raise FileNotFoundError(f"{DOCS_DIR} not found. Create it and add your .docx files.")

    new_texts = []

    for filename in os.listdir(DOCS_DIR):
        if filename.lower().endswith(".docx"):
            file_path = os.path.join(DOCS_DIR, filename)
            text = extract_text_from_docx(file_path)
            if text.strip():
                new_texts.append(text)

    if new_texts:
        # 🔑 Safely extend global context (does not overwrite CEO history)
        vector_store.add_texts("global", new_texts)
        print(f"✅ Added {len(new_texts)} Word documents to global context in vector store.")
    else:
        print("⚠️ No .docx files found or no text extracted.")

if __name__ == "__main__":
    main()
